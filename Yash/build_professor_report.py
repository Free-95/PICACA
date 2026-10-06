from pathlib import Path
import json
import pandas as pd
from docx import Document
from docx.shared import Inches, Pt, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.enum.table import WD_TABLE_ALIGNMENT, WD_CELL_VERTICAL_ALIGNMENT
from docx.enum.style import WD_STYLE_TYPE
from docx.oxml import OxmlElement
from docx.oxml.ns import qn

ROOT = Path(r"E:\ISA_PROJECT")
OUT = ROOT / "CONTEXT_RISK_V3_PROFESSOR_REPORT.docx"
DATA = ROOT / "processed" / "swat_stage1_context_risk_v3.csv"
PLOTS = ROOT / "plots" / "context_risk_v3"
CAL = ROOT / "context_engine" / "v3_artifacts" / "v3_calibration.json"

def field(paragraph, instruction):
    run = paragraph.add_run()
    begin = OxmlElement('w:fldChar'); begin.set(qn('w:fldCharType'), 'begin')
    instr = OxmlElement('w:instrText'); instr.set(qn('xml:space'), 'preserve'); instr.text = instruction
    separate = OxmlElement('w:fldChar'); separate.set(qn('w:fldCharType'), 'separate')
    text = OxmlElement('w:t'); text.text = '1'
    separate.append(text)
    end = OxmlElement('w:fldChar'); end.set(qn('w:fldCharType'), 'end')
    run._r.extend([begin, instr, separate, end])

def shade(cell, fill):
    tcPr = cell._tc.get_or_add_tcPr(); shd = OxmlElement('w:shd'); shd.set(qn('w:fill'), fill); tcPr.append(shd)

def border(cell):
    tcPr = cell._tc.get_or_add_tcPr(); b = tcPr.first_child_found_in('w:tcBorders')
    if b is None:
        b = OxmlElement('w:tcBorders'); tcPr.append(b)
    for e in ('top','left','bottom','right','insideH','insideV'):
        x = OxmlElement(f'w:{e}'); x.set(qn('w:val'),'single'); x.set(qn('w:sz'),'4'); x.set(qn('w:color'),'D9D9D9'); b.append(x)

def set_cell(cell, text, bold=False, color=None):
    cell.text = ''
    p = cell.paragraphs[0]; p.paragraph_format.space_after = Pt(0); p.paragraph_format.space_before = Pt(0)
    r = p.add_run(str(text)); r.bold = bold; r.font.size = Pt(8.5)
    if color: r.font.color.rgb = RGBColor(*color)
    cell.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER
    border(cell)

def table(doc, headers, rows, widths=None):
    t = doc.add_table(rows=1, cols=len(headers)); t.alignment = WD_TABLE_ALIGNMENT.CENTER; t.style = 'Table Grid'
    t.autofit = False
    for i,h in enumerate(headers):
        c=t.rows[0].cells[i]; shade(c,'17365D'); set_cell(c,h,True,(255,255,255))
        if widths: c.width=Inches(widths[i])
    for j,row in enumerate(rows):
        cells=t.add_row().cells
        for i,v in enumerate(row):
            if j%2: shade(cells[i],'EDF3F8')
            set_cell(cells[i],v)
            if widths: cells[i].width=Inches(widths[i])
    doc.add_paragraph().paragraph_format.space_after=Pt(3)
    return t

def add_heading(doc, text, level=1):
    p=doc.add_heading(text, level=level); p.paragraph_format.space_before=Pt(12 if level==1 else 7); p.paragraph_format.space_after=Pt(5)
    return p

def para(doc, text, bold_lead=None):
    p=doc.add_paragraph(); p.paragraph_format.space_after=Pt(5); p.paragraph_format.line_spacing=1.12
    if bold_lead and text.startswith(bold_lead):
        p.add_run(bold_lead).bold=True; p.add_run(text[len(bold_lead):])
    else: p.add_run(text)
    return p

def caption(doc, num, title, text):
    p=doc.add_paragraph(); p.style='Caption'; p.alignment=WD_ALIGN_PARAGRAPH.CENTER
    p.add_run(f'Figure {num}. {title}. ').bold=True; p.add_run(text)

def fig(doc, path, num, title, explanation, width=6.15):
    p=doc.add_paragraph(); p.alignment=WD_ALIGN_PARAGRAPH.CENTER; p.add_run().add_picture(str(path),width=Inches(width))
    caption(doc,num,title,explanation)
    para(doc, 'Key observation: ' + explanation)

def add_bullets(doc, items):
    for x in items:
        p=doc.add_paragraph(style='List Bullet'); p.add_run(x); p.paragraph_format.space_after=Pt(2)

def main():
    df=pd.read_csv(DATA)
    df['Timestamp']=pd.to_datetime(df['Timestamp'])
    cal=json.loads(CAL.read_text())
    normal=df[df['Normal/Attack'].eq('Normal')].copy().sort_values('Timestamp')
    cut=int(len(normal)*0.8)
    cal_idx=set(normal.index[:cut]); val_idx=set(normal.index[cut:])
    df['partition']='Attack evaluation'
    df.loc[df.index.isin(cal_idx),'partition']='Normal calibration'
    df.loc[df.index.isin(val_idx),'partition']='Normal validation'
    cmd=df[df.event_type.ne('NONE')].copy()
    groups=['Normal calibration','Normal validation','Attack evaluation']
    # Actual report statistics
    cmd_counts=cmd.groupby('partition').size().reindex(groups).fillna(0).astype(int)
    phase=(df.groupby(['partition','process_phase']).size().groupby(level=0).apply(lambda s:100*s/s.sum()).round(2))
    command_counts=cmd.groupby(['partition','event_type']).size()
    means=cmd.groupby('partition')[['context_risk_max','context_risk_weighted','risk_command_sequence','risk_phase_rarity']].mean()
    allmeans=df.groupby('partition')[['context_risk_max','context_risk_weighted']].mean()

    doc=Document(); sec=doc.sections[0]; sec.top_margin=Inches(.7); sec.bottom_margin=Inches(.65); sec.left_margin=Inches(.75); sec.right_margin=Inches(.75)
    styles=doc.styles
    styles['Normal'].font.name='Aptos'; styles['Normal']._element.rPr.rFonts.set(qn('w:ascii'),'Aptos'); styles['Normal'].font.size=Pt(10)
    for name,size,color in [('Title',24,(0,0,0)),('Heading 1',15,(0,0,0)),('Heading 2',12,(0,0,0)),('Caption',8,(70,70,70))]:
        s=styles[name]; s.font.name='Aptos Display' if name!='Caption' else 'Aptos'; s._element.rPr.rFonts.set(qn('w:ascii'),s.font.name); s.font.size=Pt(size); s.font.color.rgb=RGBColor(*color)
    # Header/footer
    hdr=sec.header.paragraphs[0]; hdr.alignment=WD_ALIGN_PARAGRAPH.RIGHT; r=hdr.add_run('Context Risk V3 | SWaT Stage 1'); r.font.size=Pt(8); r.font.color.rgb=RGBColor(89,89,89)
    f=sec.footer.paragraphs[0]; f.alignment=WD_ALIGN_PARAGRAPH.CENTER; f.add_run('Professor Facing Technical Report | Page '); field(f,'PAGE')
    # cover
    doc.add_paragraph().paragraph_format.space_after=Pt(70)
    p=doc.add_paragraph(style='Title'); p.alignment=WD_ALIGN_PARAGRAPH.CENTER; p.add_run('Context Risk V3 Command Context Anomaly Model')
    p=doc.add_paragraph(); p.alignment=WD_ALIGN_PARAGRAPH.CENTER; r=p.add_run('SWaT Stage 1 Professor Facing Technical Report'); r.bold=True; r.font.size=Pt(15)
    doc.add_paragraph('\n')
    for line in ['Scope: Command frequency, repetition, timing, sequence, and phase rarity','Data source: SWaT Stage 1 Network logs, parts 01 through 08','Final deliverable: Calibrated command context risk component']:
        p=doc.add_paragraph(); p.alignment=WD_ALIGN_PARAGRAPH.CENTER; p.add_run(line).font.size=Pt(11)
    doc.add_page_break()
    # TOC
    p=doc.add_heading('Table of Contents',0); p.alignment=WD_ALIGN_PARAGRAPH.CENTER
    toc=doc.add_paragraph(); field(toc,'TOC \\o "1-2" \\h \\z \\u')
    para(doc,'The table of contents is a Word field. In Microsoft Word, update fields if page numbers are not refreshed automatically.')
    doc.add_page_break()
    # 1
    add_heading(doc,'1 Executive Summary')
    para(doc,'Context Risk V3 is a transparent command-context anomaly model for SWaT Stage 1. It answers one constrained question: whether a command is contextually unusual based on command frequency, repetition, timing, sequence, and process phase rarity. The model was calibrated only on the first 80 percent of chronological Normal data, while the remaining Normal data and Attack data were retained for evaluation.')
    para(doc,'The repository results support V3 as a contextual component, not as a standalone attack classifier. Command-level mean max and weighted risks are higher for Attack commands than for Normal validation commands, while all-row Attack means are lower because attack periods contain comparatively few command rows. This distinction is central to interpretation.')
    add_heading(doc,'2 Problem Definition and Scope of Responsibility')
    para(doc,'The V3 responsibility is limited to command-context anomaly scoring. It does not evaluate tank overflow or underflow, predict tank level, use slopes or trends as risk evidence, perform physics or dynamics reasoning, make ALLOW, DELAY_ALERT, or BLOCK_ALARM decisions, use neural networks, use attack labels inside risk calculations, or tune thresholds using attack performance.')
    para(doc,'The final deliverable is SWaT preprocessing plus event extraction plus context feature engineering plus Context Risk V3 using calibrated command scoring and smoothed Markov command-sequence modeling.', 'The final deliverable is ')
    add_heading(doc,'3 Previous Fallback Approach and Why V3')
    para(doc,'Git history records an earlier time-series context-model direction alongside the V3 addition. The project documentation also describes a prior context-engine pipeline containing physical-dynamics and trend features. These artifacts are historical context rather than part of the V3 responsibility.')
    para(doc,'The final V3 design moved to a transparent discrete command-context formulation. This better matches the stated responsibility because every score can be decomposed into command frequency, repetition, timing, sequence probability, or phase-conditioned rarity. No performance number is claimed for the previous approach because the inspected V3 artifacts do not provide a comparable verified result.')
    add_heading(doc,'4 Dataset and Reproducible Pipeline')
    para(doc,'The raw source is Network, consisting of eight sorted SWaT log partitions named part01 through part08. The inspected generated output contains 449,919 rows: 395,298 Normal and 54,621 Attack.')
    table(doc,['Stage','Repository artifact','Purpose'],[
        ['Raw input','Network/ parts 01 to 08','Source SWaT log partitions'],['Preprocessing','swat_cleaner.py','Stage 1 cleaned data'],['Event extraction','swat_event_extractor.py','Actuator transition event types'],['Context features','swat_context_features.py','Rolling command and phase context'],['V3 scoring','context_engine/context_risk_v3.py','Calibrated five-component risk scores'],['Outputs','processed/swat_stage1_context_risk_v3.csv and companion CSVs','Frozen results used in this report']], [1.0,1.85,3.75])
    add_heading(doc,'5 Command Row and Process Phase Definitions')
    para(doc,'A command row is exactly one timestamp for which event_type is not NONE. A compound event such as P101_OFF;P102_ON remains one command-row timestamp. Treating it as one timestamp preserves the real observation for frequency, repetition, timing, and sequence: one co-occurring command context rather than two artificially separated commands.')
    table(doc,['Condition','process_phase'],[['MV101_state is Transition','TRANSITIONING'],['MV101_state is Open and both pumps are OFF','FILLING'],['MV101_state is Closed and any pump is ON','DRAINING'],['MV101_state is Open and any pump is ON','TRANSFERRING'],['MV101_state is Closed and both pumps are OFF','HOLDING'],['Otherwise','UNKNOWN']],[4.7,1.9])
    para(doc,'Process phase is a categorical context for command behavior only. It is not a physical or dynamics model.')
    add_heading(doc,'6 V3 Architecture')
    para(doc,'Network logs flow through cleaning, event extraction, and context feature engineering into V3 scoring. The fitted calibration artifact records frequency percentiles, timing parameters, the command and phase vocabulary, smoothing configuration, weights, and sparse-transition audit. The model output retains component scores, two aggregate scores, levels, and a context reason.')
    add_heading(doc,'7 Five Risk Components')
    table(doc,['Component','Inputs or probability','Implemented interpretation'],[
        ['risk_cmd_frequency','command_frequency; command_frequency_60s','Rolling 300-second and 60-second command burstiness; calibrated Normal percentiles form the ramps.'],
        ['risk_cmd_repetition','command_repetition; consecutive_same_command','Repeated command type and consecutive identical-command behavior. Normal MV101 transitions can occur close together, so repetition alone is not automatically suspicious.'],
        ['risk_cmd_timing','time_since_last_command; time_since_same_command','Short inter-arrival timing is evaluated against log-normal parameters fitted on Normal calibration command rows.'],
        ['risk_command_sequence','P(current | previous); P(current | previous, phase); fallback P(current | phase)','First-order Markov contextual rarity. Laplace add-one smoothing and raw support checks temper rare transitions.'],
        ['risk_phase_rarity','P(current command | current phase)','Contextual rarity of a command in its current categorical phase, not physical safety.']], [1.45,2.25,2.9])
    para(doc,'The sequence model uses Laplace smoothing k = 1. For a previous command and phase, the smoothed probability is (count(previous, phase, current) + 1) divided by (sum of counts for that previous command and phase + vocabulary size). When a transition has raw support below 3, its sequence risk is capped at 0.80; unseen or sparse transitions are therefore not automatically assigned maximum risk.')
    para(doc,f"The calibration artifact records {cal['sparse_audit']['total_normal_cmd_rows']} Normal command rows and {cal['sparse_audit']['sparse_transition_rows']} sparse transitions ({cal['sparse_audit']['sparse_pct']} percent).")
    add_heading(doc,'8 Calibration and Leakage Prevention')
    table(doc,['Partition','Rows','Use'],[['Normal calibration','316,238','First chronological 80 percent of Normal rows; fit all models'],['Normal validation','79,060','Last chronological 20 percent of Normal rows; validation only'],['Attack evaluation','54,621','Final evaluation only']], [2.0,1.2,3.4])
    para(doc,'Chronological splitting preserves the observed order of commands and avoids mixing future operating behavior into calibration. Attack rows do not influence percentile thresholds, timing thresholds, repetition thresholds, frequency thresholds, Markov probabilities, phase rarity probabilities, weights, or threshold selection.')
    add_heading(doc,'9 Context Risk Scores and Levels')
    para(doc,'context_risk_max = max(risk_cmd_frequency, risk_cmd_repetition, risk_cmd_timing, risk_command_sequence, risk_phase_rarity).')
    para(doc,'context_risk_weighted = 0.20 frequency + 0.20 repetition + 0.20 timing + 0.25 sequence + 0.15 phase rarity.')
    para(doc,'Both scores are retained: the maximum score preserves a single severe contextual signal, while the weighted score represents the combined burden across components. The repository does not designate one as the final system score.')
    table(doc,['Range','context_level_max and context_level_weighted'],[['0.00 to below 0.30','LOW'],['0.30 to below 0.60','MEDIUM'],['0.60 to below 0.85','HIGH'],['0.85 to 1.00','VERY_HIGH']],[2.1,4.5])
    add_heading(doc,'10 Experimental Results')
    table(doc,['Partition','Rows','Command rows','Max mean','Weighted mean'],[[g,f"{int((df.partition==g).sum()):,}",f"{cmd_counts[g]:,}",f"{allmeans.loc[g,'context_risk_max']:.4f}",f"{allmeans.loc[g,'context_risk_weighted']:.4f}"] for g in groups],[1.65,1.05,1.25,1.15,1.3])
    para(doc,'For Normal validation, the all-row max and weighted means are 0.0652 and 0.0133. For Attack evaluation, they are 0.0122 and 0.0026. The generated metrics CSV reports all-Normal means rather than validation-only values; the validation values above were verified directly from the frozen risk output using the model chronological split.')
    table(doc,['Command-level measure','Normal validation','Attack evaluation'],[['Max score mean',f"{means.loc['Normal validation','context_risk_max']:.4f}",f"{means.loc['Attack evaluation','context_risk_max']:.4f}"],['Weighted score mean',f"{means.loc['Normal validation','context_risk_weighted']:.4f}",f"{means.loc['Attack evaluation','context_risk_weighted']:.4f}"],['Sequence risk mean',f"{means.loc['Normal validation','risk_command_sequence']:.4f}",f"{means.loc['Attack evaluation','risk_command_sequence']:.4f}"],['Phase rarity mean',f"{means.loc['Normal validation','risk_phase_rarity']:.4f}",f"{means.loc['Attack evaluation','risk_phase_rarity']:.4f}"],['ROC AUC','0.5695 max; 0.5615 weighted','Command rows'],['PR AUC','0.2441 max; 0.2839 weighted','Command rows']],[2.0,2.15,2.15])
    para(doc,'All-row Attack mean risk is lower because V3 is command-context focused, Attack periods contain only 39 command rows, and many Attack rows are non-command steady-state observations. All-row averaging is therefore not an appropriate standalone attack-detection measure for this model. Command-row analysis is aligned to V3: Attack command rows have higher average max, weighted, sequence, and phase-rarity scores than Normal validation command rows. V3 remains a contextual component, not a standalone attack classifier.')
    add_heading(doc,'11 Command Distribution')
    rows=[]
    types=sorted(cmd.event_type.unique())
    for et in types:
        rows.append([et]+[str(int(command_counts.get((g,et),0))) for g in groups])
    table(doc,['Command event','Normal calibration','Normal validation','Attack'],rows,[2.55,1.15,1.15,1.15])
    para(doc,'Calibration is dominated by MV101_CHANGE, P101_ON, and P101_OFF. Attack command rows include compound commands absent from the calibration vocabulary, such as P101_OFF;P102_OFF and P101_OFF;P102_ON. Such commands can be contextually unusual because their phase-conditioned and transition probabilities are low; this does not establish physical danger.')
    add_heading(doc,'12 Process Phase Distribution')
    rows=[]
    phases=['TRANSFERRING','DRAINING','FILLING','HOLDING','TRANSITIONING','UNKNOWN']
    for ph in phases:
        vals=[phase.get((g,ph),0) for g in groups]
        if any(vals): rows.append([ph]+[f'{v:.2f}%' for v in vals])
    table(doc,['Phase','Normal calibration','Normal validation','Attack'],rows,[2.15,1.55,1.55,1.15])
    para(doc,'These percentages describe process context distribution across rows in each partition. The Attack evaluation period is predominantly HOLDING, which reinforces why its all-row command-context average is low. This is an observation about contextual command availability, not a physical-safety claim.')
    add_heading(doc,'13 Graphical Analysis')
    fig(doc,PLOTS/'context_risk_distribution.png',1,'Context Risk Distribution','The all-row distributions show a heavy concentration of zero risk, especially in the Attack period, because V3 only assigns context risk when command-related behavior is present.')
    fig(doc,PLOTS/'context_risk_components.png',2,'Component Risk Timeline','The five component traces expose which command-context mechanism contributes at each command event, supporting inspection rather than opaque aggregation.')
    fig(doc,PLOTS/'command_row_threshold_curve.png',3,'Command Row Threshold Curve','The generated curve provides threshold trade-offs for the two retained scores. It is descriptive evaluation output, not a threshold tuned with Attack data.')
    fig(doc,PLOTS/'top_risky_commands.png',4,'Top Risky Command Rows','The highest-risk examples combine rare sequence or phase context with frequency, repetition, or timing contributions. The plot distinguishes their Normal and Attack labels for evaluation.')
    add_heading(doc,'14 Top Risky Command Examples')
    examples=pd.read_csv(ROOT/'processed'/'context_risk_v3_top_examples.csv').head(6)
    rows=[]
    for _,r in examples.iterrows():
        rows.append([str(r['Timestamp']),r['event_type'],r['process_phase'],f"{r['risk_command_sequence']:.2f}",f"{r['risk_phase_rarity']:.2f}",f"{r['context_risk_max']:.2f}",f"{r['context_risk_weighted']:.2f}"])
    table(doc,['Timestamp','Command','Phase','Sequence','Phase rarity','Max','Weighted'],rows,[1.15,1.15,1.05,.75,.78,.55,.7])
    para(doc,'P101_OFF;P102_OFF in HOLDING has sequence risk 0.80 and phase-rarity risk 1.00, making the compound command contextually unusual. MV101_CHANGE in FILLING has sequence and phase-rarity risks of 1.00, with max risk 1.00 and weighted risk 0.8067. P101_OFF;P102_ON in TRANSFERRING has sequence risk 0.80 and phase rarity 1.00. These are contextual explanations only; they do not label the commands physically dangerous.')
    add_heading(doc,'15 What Worked')
    add_bullets(doc,['Raw Network dataset integration, preprocessing, event extraction, and context feature engineering are present in the repository.','Calibration is reproducible through v3_calibration.json, including Normal-only parameters, smoothing, weights, and sparse-transition audit.','The model provides five auditable components, two aggregate scores, context levels, context reasons, metrics, threshold curves, plots, and top-risk examples.','The Markov model combines smoothed sequence modeling with minimum-support handling, avoiding automatic maximum risk for sparse transitions.'])
    add_heading(doc,'16 Limitations and Fallback Philosophy')
    add_bullets(doc,['Attack command rows are relatively few: 39 of 54,621 Attack rows.','Attack periods are dominated by non-command rows, so all-row risk is a poor standalone measurement of attack detection for V3.','Context risk is not physical safety risk and cannot capture physical consequences from command context alone.','Command-level discrimination is modest: ROC-AUC is 0.5695 for max and 0.5615 for weighted; PR-AUC is 0.2441 and 0.2839.','Sparse command transitions exist even in Normal behavior.'])
    para(doc,'When contextual command evidence is insufficient, V3 does not compensate by adding physics or sensor prediction. Those concerns are outside this model responsibility.')
    add_heading(doc,'17 Why V3 Was Kept')
    para(doc,'The inspected V3 report and calibration artifact document an audit with no confirmed label inversion, timestamp problem, calibration contamination, Attack leakage, command parsing problem, or confirmed implementation bug. Given that evidence and its aligned, transparent scope, V3 remains frozen as a contextual component despite modest standalone discrimination. It is not presented as a perfect detector.')
    add_heading(doc,'18 Final Deliverable and Repository Files')
    para(doc,'Final deliverable: SWaT preprocessing plus event extraction plus context feature engineering plus Context Risk V3 using calibrated command scoring and smoothed Markov command-sequence modeling.')
    table(doc,['Role','Exact file or directory'],[['Preprocessing','swat_cleaner.py'],['Event extraction','swat_event_extractor.py'],['Context features','swat_context_features.py'],['V3 model','context_engine/context_risk_v3.py'],['Calibration artifact','context_engine/v3_artifacts/v3_calibration.json'],['Primary output','processed/swat_stage1_context_risk_v3.csv'],['Evaluation outputs','processed/context_risk_v3_metrics.csv; context_risk_v3_command_metrics.csv; context_risk_v3_threshold_curve.csv; context_risk_v3_top_examples.csv'],['Graphs','plots/context_risk_v3/']], [1.65,4.9])
    add_heading(doc,'19 Conclusion')
    para(doc,'Context Risk V3 provides an interpretable, Normal-calibrated view of command-context anomaly evidence for SWaT Stage 1. Its value is in explaining whether a command is rare, repeated, unusually timed, sequence-inconsistent, or phase-rare. The model deliberately remains within that boundary and preserves its two candidate scores for downstream use without making a physical-safety or final-decision claim.')
    # update fields on open
    settings=doc.settings.element; update=OxmlElement('w:updateFields'); update.set(qn('w:val'),'true'); settings.append(update)
    doc.core_properties.title='Context Risk V3 Command Context Anomaly Model for SWaT Stage 1'
    doc.core_properties.subject='Professor Facing Technical Report'
    doc.save(OUT)

if __name__=='__main__': main()
